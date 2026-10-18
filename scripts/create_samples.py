"""Create real, fictional mixed-format fixtures, never private CVs."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import simpleSplit
from backend.data import CANDIDATES,SAMPLE_JOB

root=Path(__file__).resolve().parent.parent/'samples'
root.mkdir(exist_ok=True)
for c in CANDIDATES:
    p=root/c['filename']
    if p.suffix=='.docx':
        doc=Document();doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(10 if c['layout']=='compact' else 11)
        for i,line in enumerate(c['text'].splitlines()):
            doc.add_paragraph(line,'Title' if i==0 else 'Heading 2' if line.isupper() else 'Normal')
        doc.save(p)
    else:
        cv=Canvas(str(p),pagesize=A4);y=790
        for i,line in enumerate(c['text'].splitlines()):
            font='Times-Roman' if c['layout']=='gap' else 'Helvetica'
            size=22 if i==0 else 10
            cv.setFont(font,size)
            for part in simpleSplit(line,font,size,470):
                if y<60: cv.showPage();cv.setFont(font,size);y=790
                cv.drawString(56,y,part);y-=26 if i==0 else 17
        cv.save()
doc=Document()
for line in SAMPLE_JOB.splitlines():doc.add_paragraph(line)
doc.save(root/'Sample_Job.docx')
(root/'Sample_Job.txt').write_text(SAMPLE_JOB,encoding='utf-8')
print('Created six fictional CVs and the sample job.')
